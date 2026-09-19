import base64
import io
import os

import pymupdf
from docx import Document
from dotenv import load_dotenv
from fastapi import UploadFile
from openai import OpenAI
from pptx import Presentation
from pypdf import PdfReader

load_dotenv()

_vision_client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
_vision_model = "gpt-6-astra"  # must support image input

# If a PDF's normal text layer extracts to fewer characters than this,
# treat it as a scanned/image-based PDF and fall back to OCR via the
# vision model instead.
MIN_TEXT_LENGTH = 20


def _extract_text_from_image_bytes(image_bytes: bytes, mime_type: str) -> str:
    """Send an image to a vision-capable OpenAI model and get back any readable text."""
    b64 = base64.b64encode(image_bytes).decode("utf-8")

    response = _vision_client.responses.create(
        model=_vision_model,
        input=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": (
                            "Transcribe all readable text in this image exactly as it "
                            "appears. Return only the extracted text — no commentary, "
                            "no markdown formatting, no descriptions of the image itself."
                        ),
                    },
                    {
                        "type": "input_image",
                        "image_url": f"data:{mime_type};base64,{b64}",
                    },
                ],
            }
        ],
    )
    return response.output_text.strip()


def _extract_text_from_scanned_pdf(pdf_bytes: bytes) -> str:
    """Render each page of a scanned/image-only PDF to a PNG and OCR it via
    the vision model, then join all pages together."""
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    page_texts = []
    for page in doc:
        pix = page.get_pixmap(dpi=200)
        image_bytes = pix.tobytes("png")
        page_texts.append(_extract_text_from_image_bytes(image_bytes, "image/png"))
    doc.close()
    return "\n\n".join(page_texts)


async def extract_text_from_upload(file: UploadFile) -> str:
    """Read an uploaded file entirely in memory and pull out plain text.
    Nothing is written to disk, in keeping with the 'don't persist raw
    uploads' approach for this project.
    """
    ext = file.filename.lower().rsplit(".", 1)[-1]
    raw = await file.read()
    buffer = io.BytesIO(raw)

    if ext == "pdf":
        reader = PdfReader(buffer)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

        if len(text.strip()) < MIN_TEXT_LENGTH:
            # No real text layer found — likely a scanned PDF. Fall back to OCR.
            text = _extract_text_from_scanned_pdf(raw)

        return text

    elif ext == "docx":
        doc = Document(buffer)
        return "\n".join(p.text for p in doc.paragraphs)

    elif ext == "pptx":
        prs = Presentation(buffer)
        lines = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            lines.append(run.text)
        return "\n".join(lines)

    elif ext == "txt":
        return raw.decode("utf-8", errors="ignore")

    elif ext in ("jpg", "jpeg", "png"):
        mime_type = "image/jpeg" if ext in ("jpg", "jpeg") else "image/png"
        return _extract_text_from_image_bytes(raw, mime_type)

    else:
        raise ValueError(f"Unsupported file type: .{ext}")